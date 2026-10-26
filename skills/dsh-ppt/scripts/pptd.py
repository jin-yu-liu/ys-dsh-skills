#!/usr/bin/env python3
"""DSH PPT: list templates, start a PPTD project, check it and render an editable PPTX.

Runs on the Python returned by DSH's `load_workspace_dependencies` (python-pptx,
lxml, Pillow). PyYAML is used when installed, otherwise the vendored copy in
`_vendor/` is loaded. Every command prints JSON or plain text on stdout and never
touches files outside the paths it is given.

    python pptd.py templates
    python pptd.py show <template-id> [--lang zh|en] [--page N]
    python pptd.py new <template-id> <project-dir> [--lang zh|en] [--pages 1,4,12]
    python pptd.py inspect <project-dir> [--page N]
    python pptd.py check <project-dir>
    python pptd.py render <project-dir> -o <out.pptx> [--platform mac|win|linux]
"""
from __future__ import annotations

import argparse
import copy
import html
import json
import math
import os
import re
import shutil
import sys
import tempfile
import unicodedata
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
try:
    import yaml  # type: ignore
except ImportError:  # DSH's bundled Python has no YAML parser.
    sys.path.insert(0, str(HERE / "_vendor"))
    import yaml  # type: ignore  # noqa: E402

PT = 12700  # EMU per point
ELEMENT_TYPES = {"text", "shape", "line", "chart", "table", "image"}
SHAPE_NAMES = {
    "rect": "RECTANGLE", "roundRect": "ROUNDED_RECTANGLE", "ellipse": "OVAL",
    "triangle": "ISOSCELES_TRIANGLE", "diamond": "DIAMOND", "hexagon": "HEXAGON",
    "pentagon": "REGULAR_PENTAGON", "homePlate": "PENTAGON", "chevron": "CHEVRON",
    "rightArrow": "RIGHT_ARROW", "leftArrow": "LEFT_ARROW", "upArrow": "UP_ARROW",
    "downArrow": "DOWN_ARROW", "parallelogram": "PARALLELOGRAM", "trapezoid": "TRAPEZOID",
    "star5": "STAR_5_POINT", "donut": "DONUT", "arc": "ARC", "plus": "CROSS",
}
CHART_TYPES = {"bar", "line", "pie", "area", "scatter"}
ARROWS = {"triangle", "arrow", "stealth", "oval", "diamond"}
HEX = re.compile(r"^#?([0-9a-fA-F]{6})([0-9a-fA-F]{2})?$")
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
ESCAPED_NEWLINE = re.compile(r"\\[nr]")
TAG = re.compile(r"<(/?)([a-zA-Z]+)([^>]*)>")
KNOWN_TAGS = {"br", "b", "strong", "i", "em", "u", "span", "sup", "sub"}


# --------------------------------------------------------------------------- utils

def die(message: str, code: int = 1) -> None:
    print(json.dumps({"status": "error", "message": message}, ensure_ascii=False))
    sys.exit(code)


def load_yaml(path: Path):
    with path.open("r", encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def dump_yaml(data, path: Path) -> None:
    text = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, width=1000)
    path.write_text(text, encoding="utf-8")


def templates_root() -> Path:
    env = os.environ.get("DSH_PPT_TEMPLATES")
    candidates = [Path(env)] if env else []
    candidates += [SKILL_DIR / "assets" / "templates", SKILL_DIR.parent.parent / "ppt" / "templates"]
    for candidate in candidates:
        if candidate.is_dir():
            return candidate
    die("找不到模板目录：应位于 skills/dsh-ppt/assets/templates，或用环境变量 DSH_PPT_TEMPLATES 指定")
    raise SystemExit  # unreachable


def template_dir(template_id: str) -> Path:
    if not re.match(r"^[a-z0-9][a-z0-9-]{0,63}$", template_id or ""):
        die(f"模板 id 无效：{template_id}")
    path = templates_root() / template_id
    if not (path / "metadata.json").is_file():
        die(f"没有模板 {template_id}；先运行 `templates` 查看可用 id")
    return path


def is_cjk(ch: str) -> bool:
    return unicodedata.east_asian_width(ch) in ("W", "F")


def has_cjk(text: str) -> bool:
    return any(is_cjk(ch) for ch in text)


def char_em(ch: str) -> float:
    """Approximate advance width of one character, in em."""
    if is_cjk(ch):
        return 1.0
    if ch == " ":
        return 0.28
    if ch in "il.,:;'|!`":
        return 0.28
    if ch in "mwMW@%":
        return 0.85
    if ch.isupper() or ch.isdigit():
        return 0.64 if ch.isupper() else 0.56
    if ch in "-–—()[]{}/\\\"*+=<>?&#$^~_":
        return 0.45 if ch not in "—" else 1.0
    return 0.52


def plain_text(text: str) -> str:
    """Text without HTML tags, with <br> as newlines."""
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    return html.unescape(TAG.sub("", text))


def norm_align(props: dict) -> dict:
    """`align` may be `[horizontal, vertical]`; split it into align + verticalAlign."""
    align = props.get("align")
    if isinstance(align, (list, tuple)):
        props = dict(props)
        props["align"] = align[0] if align else None
        if len(align) > 1 and "verticalAlign" not in props:
            props["verticalAlign"] = align[1]
    return props


def norm_font(value) -> dict:
    if isinstance(value, str):
        return {"latin": value}
    return dict(value) if isinstance(value, dict) else {}


# --------------------------------------------------------------------------- project

class Project:
    """A loaded PPTD project: manifest, theme and pages with resolved paths."""

    def __init__(self, root: Path):
        self.root = root.resolve()
        manifest = self.root / "deck.pptd"
        if root.is_file() and root.suffix == ".pptd":
            manifest, self.root = root.resolve(), root.resolve().parent
        if not manifest.is_file():
            die(f"{root} 下没有 deck.pptd")
        self.manifest_path = manifest
        self.deck = load_yaml(manifest) or {}
        size = self.deck.get("size") or [960, 540]
        self.width, self.height = float(size[0]), float(size[1])
        theme = self.deck.get("theme") or {}
        self.colors = theme.get("colors") or {}
        self.text_styles = theme.get("textStyles") or {}
        self.table_styles = theme.get("tableStyles") or {}
        self.pages: list[tuple[str, dict | None, str | None]] = []
        for rel in self.deck.get("pages") or []:
            path = (self.root / str(rel)).resolve()
            if not str(path).startswith(str(self.root) + os.sep):
                self.pages.append((str(rel), None, "页面路径越出工程目录"))
                continue
            try:
                self.pages.append((str(rel), load_yaml(path) or {}, None))
            except FileNotFoundError:
                self.pages.append((str(rel), None, "页面文件不存在"))
            except yaml.YAMLError as error:
                self.pages.append((str(rel), None, f"YAML 解析失败：{error}"))

    # Theme references -----------------------------------------------------
    def color(self, value):
        """Return (rgb hex, alpha 0..1) or None for transparent / missing."""
        if value is None:
            return None
        if isinstance(value, str) and value.startswith("$"):
            value = self.colors.get(value[1:])
            if value is None:
                return None
        if not isinstance(value, str):
            return None
        if value.lower() in ("none", "transparent"):
            return None
        match = HEX.match(value.strip())
        if not match:
            return None
        alpha = int(match.group(2), 16) / 255 if match.group(2) else 1.0
        return match.group(1).upper(), alpha

    def text_props(self, element: dict) -> dict:
        content = element.get("content") or {}
        style_ref = content.get("style") or element.get("style")
        props: dict = {}
        if isinstance(style_ref, str) and style_ref.startswith("$"):
            props.update(copy.deepcopy(self.text_styles.get(style_ref[1:]) or {}))
        props.update({k: v for k, v in content.items() if k != "style"})
        return props

    def resolve_path(self, src: str) -> Path | None:
        if not isinstance(src, str) or re.match(r"^[a-z]+://", src) or os.path.isabs(src):
            return None
        path = (self.root / src).resolve()
        return path if str(path).startswith(str(self.root) + os.sep) else None


# --------------------------------------------------------------------------- text layout

def wrap_lines(text: str, width_pt: float, size: float, spacing: float = 0.0) -> int:
    """Estimated number of rendered lines for `text` in a box `width_pt` wide."""
    if width_pt <= 0:
        return 10 ** 6
    total = 0
    for paragraph in text.split("\n"):
        lines, current = 1, 0.0
        tokens = re.findall(r"[A-Za-z0-9’'.,:;!?%$()\-/+&]+|\s|.", paragraph)
        for token in tokens:
            w = sum(char_em(ch) * size + spacing for ch in token)
            if current + w <= width_pt + 0.5:
                current += w
                continue
            if token.isspace():
                continue
            if w > width_pt:  # unbreakable token longer than a line: hard-wrap it
                for ch in token:
                    cw = char_em(ch) * size + spacing
                    if current + cw > width_pt + 0.5 and current > 0:
                        lines += 1
                        current = 0.0
                    current += cw
                continue
            lines += 1
            current = w
        total += lines
    return total


def text_metrics(props: dict, bounds) -> dict:
    text = plain_text(str(props.get("text", "")))
    size = float(props.get("fontSize") or 18)
    line_height = float(props.get("lineHeight") or 1.2)
    spacing = float(props.get("letterSpacing") or 0)
    _, _, w, h = bounds
    lines = wrap_lines(text, w, size, spacing)
    needed = lines * size * line_height
    max_lines = max(1, int((h + 0.15 * size) // (size * line_height)))
    avg = 1.0 if has_cjk(text) else 0.55
    per_line = max(1, int(w / (size * avg + spacing)))
    return {"lines": lines, "needed": needed, "maxLines": max_lines, "capacity": per_line * max_lines, "chars": len(text.replace("\n", ""))}


# --------------------------------------------------------------------------- check

def issue(code: str, message: str, file: str, page: int, element: str | None = None, **extra) -> dict:
    item = {"code": code, "message": message, "file": file, "page": page}
    if element:
        item["elementId"] = element
    item.update(extra)
    return item


def template_strings(project: Project) -> set[str]:
    template = (project.deck.get("template") or {}).get("id")
    if not isinstance(template, str):
        return set()
    try:
        root = templates_root() / template
    except SystemExit:
        return set()
    found: set[str] = set()
    for page in list(root.glob("source*/pages/*.page")):
        try:
            data = load_yaml(page) or {}
        except Exception:  # noqa: BLE001 - a broken template file only disables this hint
            continue
        for element in data.get("elements") or []:
            text = ((element.get("content") or {}).get("text"))
            if isinstance(text, str) and len(plain_text(text).strip()) >= 8:
                found.add(plain_text(text).strip())
    return found


def check_color(project: Project, value, where: str, out: list, ctx: tuple) -> None:
    if value is None:
        return
    if isinstance(value, str) and value.startswith("$"):
        if value[1:] not in project.colors:
            out.append(issue("ref-unknown", f"{where} 引用了未定义的主题颜色 {value}", *ctx))
        return
    if not (isinstance(value, str) and (HEX.match(value.strip()) or value.lower() in ("none", "transparent"))):
        out.append(issue("color-invalid", f"{where} 颜色 {value!r} 应为 #RRGGBB 或 #RRGGBBAA", *ctx))


def check_fill(project: Project, fill, where: str, out: list, ctx: tuple) -> None:
    if fill is None:
        return
    if isinstance(fill, str):
        check_color(project, fill, where, out, ctx)
        return
    if not isinstance(fill, dict) or fill.get("type", "solid") not in ("solid", "none"):
        out.append(issue("fill-unsupported", f"{where} 只支持 {{type: solid, color}} 或 {{type: none}}", *ctx))
        return
    check_color(project, fill.get("color"), where, out, ctx)


def check_text(project: Project, element: dict, props: dict, errors: list, warnings: list, ctx: tuple, leftovers: set) -> None:
    text = props.get("text")
    if not isinstance(text, str):
        errors.append(issue("text-missing", "content.text 必须是字符串", *ctx))
        return
    if ESCAPED_NEWLINE.search(text) and props.get("literalEscapes") is not True:
        errors.append(issue("text-escaped-newline", "文本含字面量 \\n 或 \\r：多行请用 YAML `|-` 写实际换行；需要原样展示时设 literalEscapes: true", *ctx))
    for tag in TAG.findall(text):
        if tag[1].lower() not in KNOWN_TAGS:
            warnings.append(issue("html-tag-ignored", f"不支持的标签 <{tag[1]}> 会被忽略", *ctx))
    check_color(project, props.get("color"), "文字", errors, ctx)
    style = props.get("style") or (element.get("content") or {}).get("style")
    if isinstance(style, str) and style.startswith("$") and style[1:] not in project.text_styles:
        errors.append(issue("ref-unknown", f"未定义的文字样式 {style}", *ctx))
    size = float(props.get("fontSize") or 18)
    if size < 7.5:
        warnings.append(issue("font-too-small", f"字号 {size}pt 过小，投影难以阅读；改写或拆页而不是缩小字号", *ctx))
    if element.get("rotation") in (90, 270, -90):
        return  # vertical labels: the box is laid out before rotation; skip capacity estimate
    metrics = text_metrics(props, element["bounds"])
    if metrics["needed"] > element["bounds"][3] * 1.12 + 2 and metrics["lines"] > metrics["maxLines"]:
        errors.append(issue(
            "text-overflow",
            f"约需 {metrics['lines']} 行，框内只容 {metrics['maxLines']} 行（textCapacity≈{metrics['capacity']} 字，当前 {metrics['chars']} 字）；精简文字、加大文本框或拆页",
            *ctx, textCapacity=metrics["capacity"], chars=metrics["chars"],
        ))
    stripped = plain_text(text).strip()
    if stripped in leftovers:
        warnings.append(issue("template-text-unchanged", f"仍是模板示例文字：“{stripped[:40]}”；替换为用户材料或删除该元素", *ctx))


def check_chart(project: Project, element: dict, errors: list, warnings: list, ctx: tuple) -> None:
    data = element.get("data") or {}
    cols, rows = data.get("cols"), data.get("rows")
    if not isinstance(cols, list) or not isinstance(rows, list) or not rows:
        errors.append(issue("chart-data-invalid", "图表需要 data.cols（列名数组）和 data.rows（非空二维数组）", *ctx))
        return
    for i, row in enumerate(rows):
        if not isinstance(row, list) or len(row) != len(cols):
            errors.append(issue("chart-data-invalid", f"data.rows[{i}] 长度应与 cols 相同（{len(cols)}）", *ctx))
            return
    series = element.get("series")
    if not isinstance(series, list) or not series:
        errors.append(issue("chart-data-invalid", "图表需要至少一个 series", *ctx))
        return
    kinds = set()
    for i, s in enumerate(series):
        kind = s.get("type")
        kinds.add(kind)
        if kind not in CHART_TYPES:
            errors.append(issue("chart-type-unsupported", f"series[{i}].type 只支持 {sorted(CHART_TYPES)}", *ctx))
        for key, column in (s.get("encode") or {}).items():
            if column not in cols:
                errors.append(issue("chart-data-invalid", f"series[{i}].encode.{key} 指向不存在的列 {column!r}", *ctx))
        fill = s.get("fill")
        for color in (fill if isinstance(fill, list) else []):
            check_color(project, color, "饼图颜色", errors, ctx)
        if isinstance(fill, dict):
            check_fill(project, fill, "系列填充", errors, ctx)
        check_color(project, s.get("lineColor"), "系列线条", errors, ctx)
    if len(kinds) > 1:
        warnings.append(issue("chart-mixed-types", f"组合图 {sorted(map(str, kinds))} 会统一按首个系列类型渲染", *ctx))
    check_fill(project, element.get("fill"), "图表背景", errors, ctx)


def check_table(project: Project, element: dict, errors: list, warnings: list, ctx: tuple, leftovers: set) -> None:
    rows = element.get("rows")
    if not isinstance(rows, list) or not rows or not all(isinstance(r, list) and r for r in rows):
        errors.append(issue("table-shape-invalid", "rows 必须是非空二维数组", *ctx))
        return
    width = max(len(r) for r in rows)
    if any(len(r) != width for r in rows):
        errors.append(issue("table-shape-invalid", f"每行单元格数应一致（最多 {width} 个）", *ctx))
    style = element.get("style")
    if isinstance(style, str) and style.startswith("$") and style[1:] not in project.table_styles:
        errors.append(issue("ref-unknown", f"未定义的表格样式 {style}", *ctx))
    for key, count in (("columnWidths", width), ("rowHeights", len(rows))):
        values = element.get(key)
        if values is not None and (not isinstance(values, list) or len(values) != count):
            errors.append(issue("table-shape-invalid", f"{key} 应有 {count} 项", *ctx))
    _, _, w, h = element["bounds"]
    heights = element.get("rowHeights") or [1 / len(rows)] * len(rows)
    widths = element.get("columnWidths") or [1 / width] * width
    total_h, total_w = sum(heights) or 1, sum(widths) or 1
    base = table_style(project, element)
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            cell = cell if isinstance(cell, dict) else {"text": "" if cell is None else str(cell)}
            text = str(cell.get("text", ""))
            cctx = ctx[:2] + (f"{ctx[2]} rows[{r}][{c}]",)
            if ESCAPED_NEWLINE.search(text) and cell.get("literalEscapes") is not True:
                errors.append(issue("text-escaped-newline", "单元格含字面量 \\n：改为实际换行", *cctx))
            props = cell_props(base, r, c, cell)
            if c < len(widths) and r < len(heights):
                cw = w * widths[c] / total_w if total_w > 1.5 else w * widths[c]
                ch = h * heights[r] / total_h if total_h > 1.5 else h * heights[r]
                metrics = text_metrics({**props, "text": text}, [0, 0, max(cw - 8, 1), max(ch - 4, 1)])
                if metrics["lines"] > metrics["maxLines"] + 1:
                    warnings.append(issue("table-cell-overflow", f"单元格约需 {metrics['lines']} 行，行高只容 {metrics['maxLines']} 行；PowerPoint 会撑高表格，可能越出页面", *cctx))
            if plain_text(text).strip() in leftovers:
                warnings.append(issue("template-text-unchanged", f"仍是模板示例文字：“{text[:40]}”", *cctx))


def check_project(project: Project) -> dict:
    errors: list = []
    warnings: list = []
    manifest = project.manifest_path.name
    if project.deck.get("version") not in ("v2", 2, "2"):
        warnings.append(issue("manifest-version", "deck.pptd 的 version 应为 v2", manifest, 0))
    if not project.pages:
        errors.append(issue("manifest-invalid", "deck.pptd 的 pages 为空", manifest, 0))
    for name, value in project.colors.items():
        check_color(project, value, f"theme.colors.{name}", errors, (manifest, 0))
    leftovers = template_strings(project)
    for index, (rel, page, problem) in enumerate(project.pages, start=1):
        if page is None:
            errors.append(issue("page-invalid", problem or "页面无法读取", rel, index))
            continue
        bg = page.get("background")
        if isinstance(bg, dict) and bg.get("type") == "image":
            if project.resolve_path(bg.get("src")) is None or not project.resolve_path(bg.get("src")).is_file():
                errors.append(issue("image-missing", f"背景图 {bg.get('src')!r} 不存在或不在工程目录内", rel, index))
        elif bg is not None:
            check_fill(project, bg, "背景", errors, (rel, index))
        notes = page.get("notes")
        if isinstance(notes, str) and ("DSH-authored" in notes or "illustrative" in notes.lower()):
            warnings.append(issue("notes-template", "演讲备注仍是模板说明；改写为本页讲稿或删除", rel, index))
        seen: set[str] = set()
        text_boxes = []
        for position, element in enumerate(page.get("elements") or []):
            eid = str(element.get("elementId") or f"#{position}")
            ctx = (rel, index, eid)
            if eid in seen:
                errors.append(issue("element-id-duplicate", "elementId 在页内重复", *ctx))
            seen.add(eid)
            kind = element.get("elementType")
            if kind not in ELEMENT_TYPES:
                errors.append(issue("element-type-unknown", f"elementType 只支持 {sorted(ELEMENT_TYPES)}", *ctx))
                continue
            bounds = element.get("bounds")
            if not (isinstance(bounds, list) and len(bounds) == 4 and all(isinstance(v, (int, float)) for v in bounds)) or bounds[2] < 0 or bounds[3] < 0:
                errors.append(issue("bounds-invalid", "bounds 应为 [x, y, width, height]（点，宽高非负）", *ctx))
                continue
            x, y, w, h = bounds
            if x < -2 or y < -2 or x + w > project.width + 2 or y + h > project.height + 2:
                errors.append(issue("out-of-page", f"元素越出页面（页面 {project.width:g}×{project.height:g}）", *ctx))
            if kind == "text" or (kind == "shape" and element.get("content")):
                props = project.text_props(element)
                check_text(project, element, props, errors, warnings, ctx, leftovers)
                if kind == "text" and str(props.get("text", "")).strip():
                    text_boxes.append((eid, bounds))
            if kind == "shape":
                if element.get("shapeName", "rect") not in SHAPE_NAMES:
                    errors.append(issue("shape-unsupported", f"shapeName 只支持 {sorted(SHAPE_NAMES)}", *ctx))
                check_fill(project, element.get("fill"), "填充", errors, ctx)
                check_color(project, (element.get("border") or {}).get("color"), "边框", errors, ctx)
            elif kind == "line":
                if len(line_points(element)) < 2:
                    errors.append(issue("line-invalid", "line 需要 points（如 '0,0 100,0'）和 viewBox", *ctx))
                check_color(project, (element.get("border") or {}).get("color"), "线条", errors, ctx)
                for end in element.get("arrow") or []:
                    if end is not None and end not in ARROWS:
                        errors.append(issue("arrow-unsupported", f"arrow 只支持 {sorted(ARROWS)} 或 null", *ctx))
            elif kind == "chart":
                check_chart(project, element, errors, warnings, ctx)
            elif kind == "table":
                check_table(project, element, errors, warnings, ctx, leftovers)
            elif kind == "image":
                path = project.resolve_path(element.get("src"))
                if path is None:
                    errors.append(issue("image-outside-project", "图片 src 必须是工程目录内的相对路径（不能用 URL、绝对路径或 ..）", *ctx))
                elif not path.is_file() or path.suffix.lower() not in IMAGE_EXT:
                    errors.append(issue("image-missing", f"图片 {element.get('src')} 不存在或格式不支持（{sorted(IMAGE_EXT)}）", *ctx))
        for i in range(len(text_boxes)):
            for j in range(i + 1, len(text_boxes)):
                a, b = text_boxes[i][1], text_boxes[j][1]
                ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
                oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                if ox > 4 and oy > 4 and ox * oy > 0.2 * min(a[2] * a[3], b[2] * b[3]):
                    warnings.append(issue("text-overlap", f"与 {text_boxes[j][0]} 的文本框重叠", rel, index, text_boxes[i][0]))
    return {
        "status": "needs_revision" if errors else "ok",
        "project": str(project.root),
        "pages": len(project.pages),
        "errors": errors,
        "warnings": warnings,
        "hint": "按 file + page + elementId 定位修改后重新 check；warnings 不阻止导出，但 template-text-unchanged / notes-template 通常意味着示例内容没替换。" if (errors or warnings) else None,
    }


# --------------------------------------------------------------------------- render helpers

def line_points(element: dict) -> list[tuple[float, float]]:
    points = element.get("points")
    out: list[tuple[float, float]] = []
    if isinstance(points, str):
        for pair in points.split():
            try:
                x, y = pair.split(",")
                out.append((float(x), float(y)))
            except ValueError:
                return []
    elif isinstance(points, list):
        for pair in points:
            if isinstance(pair, (list, tuple)) and len(pair) == 2:
                out.append((float(pair[0]), float(pair[1])))
    return out


def table_style(project: Project, element: dict) -> dict:
    style = element.get("style")
    if isinstance(style, str) and style.startswith("$"):
        return project.table_styles.get(style[1:]) or {}
    return style if isinstance(style, dict) else {}


def cell_props(style: dict, r: int, c: int, cell: dict) -> dict:
    props = dict(style.get("cellStyle") or {})
    if c == 0 and style.get("firstColStyle"):
        props.update(style["firstColStyle"])
    if r == 0 and style.get("firstRowStyle"):
        props.update(style["firstRowStyle"])
    props.update({k: v for k, v in cell.items() if k != "text"})
    return props


class Renderer:
    def __init__(self, project: Project, platform: str):
        from pptx import Presentation
        from pptx.util import Emu

        self.project = project
        self.platform = platform
        self.prs = Presentation()
        self.prs.slide_width = Emu(int(project.width * PT))
        self.prs.slide_height = Emu(int(project.height * PT))
        self.blank = self.prs.slide_layouts[6]

    # Fonts -----------------------------------------------------------------
    def fonts(self, family) -> tuple[str | None, str | None]:
        spec = norm_font(family)
        latin = spec.get("latin")
        key = {"mac": "mac", "win": "win"}.get(self.platform)
        ea = spec.get(key) if key else None
        ea = ea or spec.get("ea") or latin
        return latin, ea

    def apply_font(self, run, props: dict, overrides: dict | None = None) -> None:
        from lxml import etree
        from pptx.oxml.ns import qn
        from pptx.util import Pt

        o = overrides or {}
        font = run.font
        size = o.get("fontSize", props.get("fontSize"))
        if size:
            font.size = Pt(float(size))
        font.bold = bool(o.get("bold", props.get("bold", False)))
        font.italic = bool(o.get("italic", props.get("italic", False)))
        if o.get("underline", props.get("underline")):
            font.underline = True
        color = self.project.color(o.get("color", props.get("color")))
        if color:
            font.color.rgb = rgb(color[0])
            if color[1] < 1:
                set_alpha(run._r.get_or_add_rPr().find(qn("a:solidFill"))[0], color[1])
        latin, ea = self.fonts(props.get("fontFamily"))
        rpr = run._r.get_or_add_rPr()
        if props.get("letterSpacing"):
            rpr.set("spc", str(int(round(float(props["letterSpacing"]) * 100))))
        if o.get("baseline"):
            rpr.set("baseline", o["baseline"])
        for tag, face in (("a:latin", latin), ("a:ea", ea), ("a:cs", latin)):
            if not face:
                continue
            node = rpr.find(qn(tag))
            if node is None:
                node = etree.SubElement(rpr, qn(tag))
            node.set("typeface", face)
        # Schema order inside rPr: ln, fill, effect, highlight, uLnTx.., latin, ea, cs, sym, hlink
        fill = rpr.find(qn("a:solidFill"))
        if fill is not None:
            rpr.remove(fill)
            rpr.insert(0, fill)

    def fill_text(self, frame, props: dict, bounds) -> None:
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Pt

        props = norm_align(props)
        frame.word_wrap = True
        frame.auto_size = None
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(frame, side, Pt(float(props.get("padding", 0))))
        anchor = {"middle": MSO_ANCHOR.MIDDLE, "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(props.get("verticalAlign"), MSO_ANCHOR.TOP)
        frame.vertical_anchor = anchor
        align = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}.get(props.get("align"), PP_ALIGN.LEFT)
        size = float(props.get("fontSize") or 18)
        line_height = float(props.get("lineHeight") or 1.2)
        text = str(props.get("text", ""))
        if props.get("literalEscapes") is not True:
            text = text.replace("\r\n", "\n")
        paragraphs = parse_rich(text)
        for index, runs in enumerate(paragraphs):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.alignment = align
            para.line_spacing = Pt(size * line_height)
            if props.get("paragraphSpacing"):
                para.space_after = Pt(float(props["paragraphSpacing"]))
            if not runs:
                runs = [("", {})]
            for chunk, style in runs:
                run = para.add_run()
                run.text = chunk
                self.apply_font(run, props, style)
            # An empty paragraph still needs the font size so blank lines keep their height.
            end = para._p.get_or_add_endParaRPr()
            end.set("sz", str(int(size * 100)))

    # Elements --------------------------------------------------------------
    def frame(self, shape, element: dict) -> None:
        if element.get("rotation"):
            shape.rotation = float(element["rotation"])

    def set_fill(self, fill_format, fill) -> None:
        if isinstance(fill, str):
            fill = {"type": "solid", "color": fill}
        color = self.project.color((fill or {}).get("color")) if isinstance(fill, dict) and fill.get("type", "solid") == "solid" else None
        if color is None or color[1] == 0:
            fill_format.background()
            return
        fill_format.solid()
        fill_format.fore_color.rgb = rgb(color[0])
        if color[1] < 1:
            set_alpha(fill_format._xPr.find(qn_("a:solidFill"))[0], color[1])

    def set_line(self, line, border) -> None:
        from pptx.enum.dml import MSO_LINE
        from pptx.util import Pt

        border = border or {}
        color = self.project.color(border.get("color"))
        width = float(border.get("width", 0) or 0)
        if width <= 0 or color is None or color[1] == 0:
            line.fill.background()
            return
        line.color.rgb = rgb(color[0])
        line.width = Pt(width)
        dash = {"dash": MSO_LINE.DASH, "dot": MSO_LINE.ROUND_DOT, "dashDot": MSO_LINE.DASH_DOT, "longDash": MSO_LINE.LONG_DASH}.get(border.get("style"))
        if dash:
            line.dash_style = dash
        if color[1] < 1:
            set_alpha(line._get_or_add_ln().find(qn_("a:solidFill"))[0], color[1])

    def add_text(self, slide, element: dict) -> None:
        x, y, w, h = element["bounds"]
        box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        box.name = str(element.get("elementId", "text"))
        self.fill_text(box.text_frame, self.project.text_props(element), element["bounds"])
        self.frame(box, element)

    def add_shape(self, slide, element: dict) -> None:
        from pptx.enum.shapes import MSO_SHAPE

        x, y, w, h = element["bounds"]
        kind = getattr(MSO_SHAPE, SHAPE_NAMES.get(element.get("shapeName", "rect"), "RECTANGLE"))
        shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(max(w, 0.01)), emu(max(h, 0.01)))
        shape.name = str(element.get("elementId", "shape"))
        shape.shadow.inherit = False
        if element.get("shapeName") == "roundRect" and element.get("radius") is not None and min(w, h) > 0:
            shape.adjustments[0] = max(0.0, min(0.5, float(element["radius"]) / min(w, h)))
        self.set_fill(shape.fill, element.get("fill"))
        self.set_line(shape.line, element.get("border"))
        if element.get("content"):
            props = self.project.text_props(element)
            props.setdefault("padding", 6)
            props.setdefault("verticalAlign", "middle")
            self.fill_text(shape.text_frame, props, element["bounds"])
        else:
            shape.text_frame.text = ""
        self.frame(shape, element)

    def add_line(self, slide, element: dict) -> None:
        from pptx.enum.shapes import MSO_CONNECTOR

        x, y, w, h = element["bounds"]
        points = line_points(element)
        view = element.get("viewBox") or [w or 1, h or 1]
        sx = w / float(view[0]) if view[0] else 1.0
        sy = h / float(view[1]) if view[1] else 1.0
        absolute = [(x + px * sx, y + py * sy) for px, py in points]
        if len(absolute) == 2:
            (x1, y1), (x2, y2) = absolute
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        else:
            builder = slide.shapes.build_freeform(emu(absolute[0][0]), emu(absolute[0][1]), scale=1.0)
            builder.add_line_segments([(emu(px), emu(py)) for px, py in absolute[1:]], close=False)
            shape = builder.convert_to_shape()
            shape.fill.background()
        shape.name = str(element.get("elementId", "line"))
        border = dict(element.get("border") or {})
        border.setdefault("width", 1)
        self.set_line(shape.line, border)
        arrows = element.get("arrow") or []
        if arrows:
            ln = shape.line._get_or_add_ln()
            from lxml import etree

            for tag, end in (("a:headEnd", arrows[0] if len(arrows) > 0 else None), ("a:tailEnd", arrows[1] if len(arrows) > 1 else None)):
                if end:
                    node = etree.SubElement(ln, qn_(tag))
                    node.set("type", end)
                    node.set("w", "med")
                    node.set("len", "med")

    def add_image(self, slide, element: dict) -> None:
        from PIL import Image

        x, y, w, h = element["bounds"]
        path = self.project.resolve_path(element["src"])
        with Image.open(path) as image:
            iw, ih = image.size
        fit = element.get("fit", "cover")
        if fit == "contain":
            scale = min(w / iw, h / ih)
            dw, dh = iw * scale, ih * scale
            pic = slide.shapes.add_picture(str(path), emu(x + (w - dw) / 2), emu(y + (h - dh) / 2), emu(dw), emu(dh))
        else:
            pic = slide.shapes.add_picture(str(path), emu(x), emu(y), emu(w), emu(h))
            if fit == "cover" and iw and ih:
                box_ratio, img_ratio = w / h if h else 1, iw / ih
                if img_ratio > box_ratio:  # too wide: crop left/right
                    crop = (1 - box_ratio / img_ratio) / 2
                    pic.crop_left = pic.crop_right = crop
                elif img_ratio < box_ratio:
                    crop = (1 - img_ratio / box_ratio) / 2
                    pic.crop_top = pic.crop_bottom = crop
        pic.name = str(element.get("elementId", "image"))
        self.frame(pic, element)

    def add_table(self, slide, element: dict) -> None:
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Pt

        x, y, w, h = element["bounds"]
        rows = element["rows"]
        n_rows, n_cols = len(rows), max(len(r) for r in rows)
        shape = slide.shapes.add_table(n_rows, n_cols, emu(x), emu(y), emu(w), emu(h))
        shape.name = str(element.get("elementId", "table"))
        table = shape.table
        table.first_row = table.horz_banding = False
        tbl_pr = table._tbl.tblPr
        style_id = tbl_pr.find(qn_("a:tableStyleId"))
        if style_id is not None:
            style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
        widths = element.get("columnWidths") or [1 / n_cols] * n_cols
        heights = element.get("rowHeights") or [1 / n_rows] * n_rows
        tw, th = sum(widths) or 1, sum(heights) or 1
        for i, value in enumerate(widths[:n_cols]):
            table.columns[i].width = emu(w * value / tw if tw > 1.5 or abs(tw - 1) < 0.05 else value)
        for i, value in enumerate(heights[:n_rows]):
            table.rows[i].height = emu(h * value / th if th > 1.5 or abs(th - 1) < 0.05 else value)
        style = table_style(self.project, element)
        for r, row in enumerate(rows):
            for c in range(n_cols):
                raw = row[c] if c < len(row) else ""
                cell_data = raw if isinstance(raw, dict) else {"text": "" if raw is None else str(raw)}
                props = cell_props(style, r, c, cell_data)
                props["text"] = str(cell_data.get("text", ""))
                props.setdefault("fontSize", 12)
                cell = table.cell(r, c)
                cell.margin_left = cell.margin_right = Pt(float(props.get("paddingX", 6)))
                cell.margin_top = cell.margin_bottom = Pt(float(props.get("paddingY", 3)))
                props = norm_align(props)
                cell.vertical_anchor = {"top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}.get(props.get("verticalAlign"), MSO_ANCHOR.MIDDLE)
                text_props = dict(props)
                text_props.pop("padding", None)
                frame = cell.text_frame
                self.fill_text_cell(frame, text_props)
                self.set_fill(cell.fill, props.get("fill"))
                set_cell_border(cell, self.project.color((props.get("border") or {}).get("color")), float((props.get("border") or {}).get("width", 0) or 0))
        for r, row in enumerate(rows):
            for c, raw in enumerate(row):
                if isinstance(raw, dict) and (raw.get("colspan", 1) > 1 or raw.get("rowspan", 1) > 1):
                    end = table.cell(min(n_rows - 1, r + raw.get("rowspan", 1) - 1), min(n_cols - 1, c + raw.get("colspan", 1) - 1))
                    table.cell(r, c).merge(end)

    def fill_text_cell(self, frame, props: dict) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt

        props = norm_align(props)
        frame.word_wrap = True
        align = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}.get(props.get("align"), PP_ALIGN.CENTER)
        size = float(props.get("fontSize") or 12)
        for index, runs in enumerate(parse_rich(str(props.get("text", "")))):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.alignment = align
            if props.get("lineHeight"):
                para.line_spacing = Pt(size * float(props["lineHeight"]))
            for chunk, style in runs or [("", {})]:
                run = para.add_run()
                run.text = chunk
                self.apply_font(run, props, style)

    def add_chart(self, slide, element: dict) -> None:
        from pptx.chart.data import CategoryChartData, XyChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
        from pptx.util import Pt

        x, y, w, h = element["bounds"]
        cols = element["data"]["cols"]
        rows = element["data"]["rows"]
        series = element["series"]
        first = series[0]
        kind = first.get("type", "bar")
        col = {name: i for i, name in enumerate(cols)}

        def numeric(name):
            return all(isinstance(r[col[name]], (int, float)) or r[col[name]] is None for r in rows)

        horizontal = False
        if kind == "pie":
            category = first.get("encode", {}).get("category", cols[0])
            values = [(s.get("name", s["encode"].get("value")), s["encode"].get("value", cols[-1])) for s in series[:1]]
        else:
            enc = first.get("encode", {})
            ex, ey = enc.get("x", cols[0]), enc.get("y", cols[-1])
            horizontal = kind == "bar" and numeric(ex) and not numeric(ey) or bool(element.get("horizontal"))
            category = ey if horizontal else ex
            values = []
            for s in series:
                e = s.get("encode", {})
                values.append((s.get("name") or (e.get("x") if horizontal else e.get("y")), e.get("x") if horizontal else e.get("y", cols[-1])))
        stacked = bool(element.get("stacked"))
        if kind == "scatter":
            data = XyChartData()
            for name, vcol in values:
                ser = data.add_series(str(name))
                for r in rows:
                    ser.add_data_point(r[col[category]], r[col[vcol]])
            chart_type = XL_CHART_TYPE.XY_SCATTER
        else:
            data = CategoryChartData()
            data.categories = [str(r[col[category]]) for r in rows]
            for name, vcol in values:
                data.add_series(str(name), [r[col[vcol]] for r in rows])
            if kind == "pie":
                chart_type = XL_CHART_TYPE.DOUGHNUT if float(first.get("innerRadius") or 0) > 0 else XL_CHART_TYPE.PIE
            elif kind == "line":
                chart_type = XL_CHART_TYPE.LINE_MARKERS if any(s.get("marker") for s in series) else XL_CHART_TYPE.LINE
            elif kind == "area":
                chart_type = XL_CHART_TYPE.AREA_STACKED if stacked else XL_CHART_TYPE.AREA
            elif horizontal:
                chart_type = XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED
            else:
                chart_type = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
        frame = slide.shapes.add_chart(chart_type, emu(x), emu(y), emu(w), emu(h), data)
        frame.name = str(element.get("elementId", "chart"))
        chart = frame.chart
        sample = " ".join(str(r[col[category]]) for r in rows) + " ".join(str(v[0]) for v in values)
        latin, ea = self.fonts(element.get("fontFamily") or {"latin": "Arial"})
        chart.font.name = ea if has_cjk(sample) else latin
        chart.font.size = Pt(float(element.get("fontSize", 10)))
        legend = element.get("legend") or {}
        chart.has_legend = bool(legend.get("show", len(series) > 1 or kind == "pie"))
        if chart.has_legend:
            chart.legend.include_in_layout = False
            chart.legend.position = {"top": XL_LEGEND_POSITION.TOP, "left": XL_LEGEND_POSITION.LEFT, "right": XL_LEGEND_POSITION.RIGHT}.get(legend.get("position"), XL_LEGEND_POSITION.BOTTOM)
        chart.has_title = False
        plot = chart.plots[0]
        if kind == "pie" and chart_type == XL_CHART_TYPE.DOUGHNUT:
            hole = plot._element.find(qn_("c:holeSize"))
            if hole is not None:
                hole.set("val", str(int(max(10, min(90, float(first["innerRadius"]) * 100)))))
        if kind in ("bar",):
            plot.gap_width = int(element.get("gapWidth", 60))
            if stacked:
                plot.overlap = 100
        show_labels = any((s.get("dataLabels") or {}).get("show") for s in series)
        if show_labels:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.show_value = True
            labels.font.size = Pt(float((first.get("dataLabels") or {}).get("fontSize", 9)))
            fmt = (first.get("dataLabels") or {}).get("format")
            if fmt:
                labels.number_format = fmt
                labels.number_format_is_linked = False
            color = self.project.color((first.get("dataLabels") or {}).get("color"))
            if color:
                labels.font.color.rgb = rgb(color[0])
        for s_def, ser in zip(series, plot.series):
            if kind == "pie":
                colors = s_def.get("fill") if isinstance(s_def.get("fill"), list) else []
                for i, point_color in enumerate(colors[: len(rows)]):
                    point = ser.points[i]
                    self.set_fill(point.format.fill, {"type": "solid", "color": point_color})
                    point.format.line.fill.background()
                continue
            if kind == "line" or kind == "scatter":
                color = self.project.color(s_def.get("lineColor") or (s_def.get("fill") or {}).get("color") if isinstance(s_def.get("fill"), dict) else s_def.get("lineColor"))
                if color:
                    ser.format.line.color.rgb = rgb(color[0])
                ser.format.line.width = Pt(float(s_def.get("width", 2)))
                ser.smooth = bool(s_def.get("smooth", False))
                marker = s_def.get("marker")
                if marker:
                    ser.marker.style = {"circle": XL_MARKER_STYLE.CIRCLE, "square": XL_MARKER_STYLE.SQUARE, "diamond": XL_MARKER_STYLE.DIAMOND, "triangle": XL_MARKER_STYLE.TRIANGLE}.get(marker.get("shape"), XL_MARKER_STYLE.CIRCLE)
                    ser.marker.size = int(marker.get("size", 6))
                    self.set_fill(ser.marker.format.fill, {"type": "solid", "color": marker.get("fill") or s_def.get("lineColor")})
                    mline = self.project.color(marker.get("lineColor"))
                    if mline:
                        ser.marker.format.line.color.rgb = rgb(mline[0])
                elif chart_type == XL_CHART_TYPE.LINE:
                    ser.marker.style = XL_MARKER_STYLE.NONE
            else:
                self.set_fill(ser.format.fill, s_def.get("fill") or {"type": "solid", "color": "#4176E6"})
                ser.format.line.fill.background()
                point_colors = s_def.get("pointColors")
                for i, point_color in enumerate(point_colors or []):
                    self.set_fill(ser.points[i].format.fill, {"type": "solid", "color": point_color})
        if kind != "pie":
            cat_axis, val_axis = chart.category_axis, chart.value_axis
            x_def, y_def = element.get("xAxis") or {}, element.get("yAxis") or {}
            cat_def, val_def = (y_def, x_def) if horizontal else (x_def, y_def)
            if horizontal:
                cat_axis.reverse_order = True  # first data row on top, as written
            for axis, spec in ((cat_axis, cat_def), (val_axis, val_def)):
                self.style_axis(axis, spec)
            for key, attr in (("min", "minimum_scale"), ("max", "maximum_scale"), ("majorUnit", "major_unit")):
                if val_def.get(key) is not None:
                    setattr(val_axis, attr, float(val_def[key]))
            if val_def.get("format"):
                val_axis.tick_labels.number_format = val_def["format"]
                val_axis.tick_labels.number_format_is_linked = False
        bg = element.get("fill")
        set_chart_area(chart, self.project.color((bg or {}).get("color") if isinstance(bg, dict) else bg), self.project.color((element.get("border") or {}).get("color")), float((element.get("border") or {}).get("width", 0) or 0))
        layout = element.get("plotLayout")
        if isinstance(layout, dict):
            set_plot_layout(chart, layout)

    def style_axis(self, axis, spec: dict) -> None:
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        from pptx.enum.dml import MSO_LINE
        from pptx.util import Pt

        if spec.get("hidden"):
            axis.visible = False
            axis.has_major_gridlines = False
            return
        label = spec.get("label")
        if isinstance(label, dict):
            if label.get("fontSize"):
                axis.tick_labels.font.size = Pt(float(label["fontSize"]))
            color = self.project.color(label.get("color"))
            if color:
                axis.tick_labels.font.color.rgb = rgb(color[0])
            position = {"none": XL_TICK_LABEL_POSITION.NONE, "low": XL_TICK_LABEL_POSITION.LOW, "high": XL_TICK_LABEL_POSITION.HIGH, "nextTo": XL_TICK_LABEL_POSITION.NEXT_TO_AXIS}.get(label.get("position"))
            if position is not None:
                axis.tick_label_position = position
        grid = spec.get("gridLine")
        if grid is False or grid is None:
            axis.has_major_gridlines = False
        elif isinstance(grid, dict):
            axis.has_major_gridlines = True
            line = axis.major_gridlines.format.line
            color = self.project.color(grid.get("color"))
            if color:
                line.color.rgb = rgb(color[0])
            line.width = Pt(float(grid.get("width", 0.5)))
            if grid.get("style") in ("dot", "dash"):
                line.dash_style = MSO_LINE.ROUND_DOT if grid["style"] == "dot" else MSO_LINE.DASH
        line_spec = spec.get("line")
        if line_spec is False:
            axis.format.line.fill.background()
        elif isinstance(line_spec, dict):
            color = self.project.color(line_spec.get("color"))
            if color:
                axis.format.line.color.rgb = rgb(color[0])
            axis.format.line.width = Pt(float(line_spec.get("width", 0.75)))

    # Slides ----------------------------------------------------------------
    def render(self, out: Path) -> int:
        adders = {"text": self.add_text, "shape": self.add_shape, "line": self.add_line, "chart": self.add_chart, "table": self.add_table, "image": self.add_image}
        count = 0
        for _rel, page, _problem in self.project.pages:
            slide = self.prs.slides.add_slide(self.blank)
            count += 1
            bg = page.get("background")
            if isinstance(bg, dict) and bg.get("type") == "image":
                path = self.project.resolve_path(bg["src"])
                pic = slide.shapes.add_picture(str(path), 0, 0, self.prs.slide_width, self.prs.slide_height)
                pic.name = "background"
            elif bg is not None:
                self.set_fill(slide.background.fill, bg)
            for element in page.get("elements") or []:
                try:
                    adders[element["elementType"]](slide, element)
                except Exception as error:  # noqa: BLE001
                    raise RuntimeError(f"{_rel} 第 {count} 页元素 {element.get('elementId')}：{type(error).__name__}: {error}") from error
            notes = page.get("notes")
            if isinstance(notes, str) and notes.strip():
                slide.notes_slide.notes_text_frame.text = notes.strip()
        title = self.project.deck.get("title")
        if isinstance(title, str):
            self.prs.core_properties.title = title
        self.prs.core_properties.author = "DSH"
        fd, tmp = tempfile.mkstemp(suffix=".pptx", dir=str(out.parent))
        os.close(fd)
        try:
            self.prs.save(tmp)
            os.replace(tmp, out)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        return count


def qn_(tag: str) -> str:
    from pptx.oxml.ns import qn

    return qn(tag)


def emu(points: float):
    from pptx.util import Emu

    return Emu(int(round(float(points) * PT)))


def rgb(hex6: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex6)


def set_alpha(color_node, alpha: float) -> None:
    from lxml import etree

    for old in color_node.findall(qn_("a:alpha")):
        color_node.remove(old)
    node = etree.SubElement(color_node, qn_("a:alpha"))
    node.set("val", str(int(round(alpha * 100000))))


def set_cell_border(cell, color, width: float) -> None:
    from lxml import etree

    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tc_pr.findall(qn_(tag)):
            tc_pr.remove(old)
    nodes = []
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.Element(qn_(tag))
        if color is None or width <= 0:
            etree.SubElement(ln, qn_("a:noFill"))
            ln.set("w", "0")
        else:
            ln.set("w", str(int(width * PT)))
            fill = etree.SubElement(ln, qn_("a:solidFill"))
            clr = etree.SubElement(fill, qn_("a:srgbClr"))
            clr.set("val", color[0])
        nodes.append(ln)
    # Borders precede the cell fill in CT_TableCellProperties.
    for offset, ln in enumerate(nodes):
        tc_pr.insert(offset, ln)


def set_chart_area(chart, fill, border, width: float) -> None:
    from lxml import etree

    def sp_pr(parent, color, line_color=None, line_width=0.0):
        for old in parent.findall(qn_("c:spPr")):
            parent.remove(old)
        node = etree.Element(qn_("c:spPr"))
        if color is None or color[1] == 0:
            etree.SubElement(node, qn_("a:noFill"))
        else:
            solid = etree.SubElement(node, qn_("a:solidFill"))
            clr = etree.SubElement(solid, qn_("a:srgbClr"))
            clr.set("val", color[0])
            if color[1] < 1:
                set_alpha(clr, color[1])
        ln = etree.SubElement(node, qn_("a:ln"))
        if line_color is None or line_color[1] == 0 or line_width <= 0:
            etree.SubElement(ln, qn_("a:noFill"))
        else:
            ln.set("w", str(int(line_width * PT)))
            solid = etree.SubElement(ln, qn_("a:solidFill"))
            etree.SubElement(solid, qn_("a:srgbClr")).set("val", line_color[0])
        return node

    space = chart._chartSpace
    # chartSpace/spPr sits after c:chart; plotArea/spPr after the axes.
    node = sp_pr(space, fill, border, width)
    chart_el = space.find(qn_("c:chart"))
    chart_el.addnext(node)
    plot_area = chart_el.find(qn_("c:plotArea"))
    plot_node = sp_pr(plot_area, None)
    ext = plot_area.find(qn_("c:extLst"))
    if ext is not None:
        ext.addprevious(plot_node)
    else:
        plot_area.append(plot_node)


def set_plot_layout(chart, layout: dict) -> None:
    from lxml import etree

    plot_area = chart._chartSpace.find(qn_("c:chart")).find(qn_("c:plotArea"))
    for old in plot_area.findall(qn_("c:layout")):
        plot_area.remove(old)
    node = etree.Element(qn_("c:layout"))
    manual = etree.SubElement(node, qn_("c:manualLayout"))
    etree.SubElement(manual, qn_("c:layoutTarget")).set("val", "inner")
    etree.SubElement(manual, qn_("c:xMode")).set("val", "edge")
    etree.SubElement(manual, qn_("c:yMode")).set("val", "edge")
    for key in ("x", "y", "w", "h"):
        etree.SubElement(manual, qn_(f"c:{key}")).set("val", str(float(layout.get(key, 0))))
    plot_area.insert(0, node)


def parse_rich(text: str) -> list[list[tuple[str, dict]]]:
    """Split light HTML into paragraphs of (text, style) runs."""
    paragraphs: list[list[tuple[str, dict]]] = [[]]
    stack: list[tuple[str, dict]] = []
    pos = 0

    def current() -> dict:
        merged: dict = {}
        for _tag, style in stack:
            merged.update(style)
        return merged

    def emit(chunk: str) -> None:
        if not chunk:
            return
        chunk = html.unescape(chunk)
        parts = chunk.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                paragraphs.append([])
            if part:
                paragraphs[-1].append((part, current()))

    for match in TAG.finditer(text):
        emit(text[pos:match.start()])
        pos = match.end()
        closing, name, attrs = match.group(1), match.group(2).lower(), match.group(3)
        if name == "br":
            paragraphs.append([])
            continue
        if closing:
            for i in range(len(stack) - 1, -1, -1):
                if stack[i][0] == name:
                    del stack[i:]
                    break
            continue
        style: dict = {}
        if name in ("b", "strong"):
            style["bold"] = True
        elif name in ("i", "em"):
            style["italic"] = True
        elif name == "u":
            style["underline"] = True
        elif name == "sup":
            style["baseline"] = "30000"
        elif name == "sub":
            style["baseline"] = "-25000"
        elif name == "span":
            css = re.search(r'style\s*=\s*["\']([^"\']*)["\']', attrs)
            for decl in (css.group(1).split(";") if css else []):
                if ":" not in decl:
                    continue
                key, value = (s.strip().lower() for s in decl.split(":", 1))
                if key == "color":
                    style["color"] = value if value.startswith("#") else f"#{value}"
                elif key == "font-size":
                    number = re.match(r"([\d.]+)", value)
                    if number:
                        style["fontSize"] = float(number.group(1))
                elif key == "font-weight":
                    style["bold"] = value in ("bold", "600", "700", "800", "900")
                elif key == "font-style":
                    style["italic"] = value == "italic"
        if name in KNOWN_TAGS:
            stack.append((name, style))
    emit(text[pos:])
    return paragraphs


# --------------------------------------------------------------------------- commands

def cmd_templates(args) -> None:
    root = templates_root()
    rows = []
    for meta_path in sorted(root.glob("*/metadata.json")):
        meta = json.loads(meta_path.read_text("utf-8"))
        d = meta.get("definition", {})
        names = d.get("localizedNames") or {}
        rows.append({
            "id": d.get("id", meta_path.parent.name),
            "name": names.get("zh") or d.get("name"),
            "nameEn": names.get("en") or d.get("name"),
            "category": d.get("category"),
            "pages": d.get("referencePageCount"),
            "density": (meta.get("semantics", {}).get("source") or {}).get("recommendedDensity"),
            "palette": d.get("palette"),
        })
    catalog = SKILL_DIR.parent.parent / "ppt" / "catalog.json"
    if catalog.is_file():
        info = {t["id"]: t for t in json.loads(catalog.read_text("utf-8")).get("templates", [])}
        for row in rows:
            if row["id"] in info:
                row["name"] = info[row["id"]]["name"]["zh"]
                row["description"] = info[row["id"]].get("description")
    if args.json:
        print(json.dumps({"templatesDir": str(root), "templates": rows}, ensure_ascii=False, indent=1))
        return
    print(f"模板目录：{root}")
    for row in rows:
        desc = row.get("description") or ""
        print(f"- {row['id']}  {row['name']} / {row['nameEn']}  [{row['category']}, {row['pages']} 页]  {desc}")


def cmd_show(args) -> None:
    root = template_dir(args.template)
    source = root / ("source-zh" if args.lang == "zh" else "source")
    if args.page:
        page = source / "pages" / f"{int(args.page):02d}.page"
        if not page.is_file():
            die(f"没有第 {args.page} 页")
        print(f"# {page}\n")
        print(page.read_text("utf-8"))
        return
    meta = json.loads((root / "metadata.json").read_text("utf-8"))
    d = meta["definition"]
    print(f"# {d['id']} — {(d.get('localizedNames') or {}).get('zh', d['name'])} / {d['name']}")
    print(f"源工程：{source}")
    print(f"设计说明：{root / 'design.md'}")
    print(f"字体：{json.dumps(d.get('fonts', {}).get('fallbacks', {}), ensure_ascii=False)}")
    print(f"配色：{json.dumps(d.get('palette', {}), ensure_ascii=False)}")
    deck = load_yaml(source / "deck.pptd")
    if deck.get("theme"):
        print(f"主题引用：colors={list((deck['theme'].get('colors') or {}).keys())} textStyles={list((deck['theme'].get('textStyles') or {}).keys())} tableStyles={list((deck['theme'].get('tableStyles') or {}).keys())}")
    design = (root / "design.md").read_text("utf-8")
    rules = re.search(r"## Expanded composition rules\s*\n(.*?)(?=\n## |\Z)", design, re.S)
    if rules:
        print("\n## 版式（来自 design.md）\n" + rules.group(1).strip()[:2400])
    print("\n## 页面概览")
    for i, rel in enumerate(deck.get("pages") or [], start=1):
        page = load_yaml(source / rel) or {}
        elements = page.get("elements") or []
        texts = sorted(
            (e for e in elements if e.get("elementType") == "text"),
            key=lambda e: -float((e.get("content") or {}).get("fontSize") or 0),
        )
        headline = plain_text(str((texts[0].get("content") or {}).get("text", ""))).replace("\n", " ") if texts else ""
        kinds: dict = {}
        for e in elements:
            kinds[e.get("elementType")] = kinds.get(e.get("elementType"), 0) + 1
        print(f"{i:02d}. [{page.get('pageType', '?')}] {headline[:50]}  {kinds}")
    print(f"\n完整页面：python {Path(__file__).name} show {args.template} --lang {args.lang} --page N")


def cmd_new(args) -> None:
    root = template_dir(args.template)
    source = root / ("source-zh" if args.lang == "zh" else "source")
    target = Path(args.project).resolve()
    if target.exists() and any(target.iterdir()) and not args.force:
        die(f"{target} 已存在且非空；换一个目录，或确认可覆盖后加 --force")
    deck = load_yaml(source / "deck.pptd")
    pages = list(deck.get("pages") or [])
    if args.pages:
        try:
            picks = [int(p) for p in args.pages.split(",") if p.strip()]
        except ValueError:
            die("--pages 应为逗号分隔的页码，如 1,3,4,12")
        if any(p < 1 or p > len(pages) for p in picks):
            die(f"页码超出范围 1..{len(pages)}")
        pages = [pages[p - 1] for p in picks]
    (target / "pages").mkdir(parents=True, exist_ok=True)
    new_pages = []
    for i, rel in enumerate(pages, start=1):
        name = f"pages/{i:02d}.page"
        shutil.copyfile(source / rel, target / name)
        new_pages.append(name)
    deck["pages"] = new_pages
    if args.title:
        deck["title"] = args.title
    deck.setdefault("template", {"id": args.template})
    dump_yaml(deck, target / "deck.pptd")
    print(json.dumps({
        "status": "created",
        "project": str(target),
        "template": args.template,
        "pages": [str(target / p) for p in new_pages],
        "next": "逐页替换示例文字/数字/备注为用户材料，增删页时同步 deck.pptd 的 pages；然后运行 check。",
    }, ensure_ascii=False, indent=1))


def cmd_inspect(args) -> None:
    project = Project(Path(args.project))
    for index, (rel, page, problem) in enumerate(project.pages, start=1):
        if args.page and index != args.page:
            continue
        if page is None:
            print(f"## {index:02d} {rel}  !! {problem}")
            continue
        print(f"## {index:02d} {rel}  [{page.get('pageType', '?')}]")
        for element in page.get("elements") or []:
            kind = element.get("elementType")
            b = element.get("bounds") or []
            head = f"  - {element.get('elementId')} {kind} {[round(v) for v in b]}"
            if kind in ("text",) or (kind == "shape" and element.get("content")):
                props = project.text_props(element)
                text = plain_text(str(props.get("text", ""))).replace("\n", "⏎")
                try:
                    cap = text_metrics(props, b)["capacity"]
                except Exception:  # noqa: BLE001
                    cap = "?"
                print(f"{head} {props.get('fontSize', '?')}pt cap≈{cap}: {text[:70]}")
            elif kind == "chart":
                print(f"{head} {[s.get('type') for s in element.get('series') or []]} cols={element.get('data', {}).get('cols')} rows={len(element.get('data', {}).get('rows') or [])}")
            elif kind == "table":
                rows = element.get("rows") or []
                print(f"{head} {len(rows)}×{max((len(r) for r in rows), default=0)}")
            elif kind == "image":
                print(f"{head} src={element.get('src')}")
            elif args.all:
                print(head)


def cmd_check(args) -> None:
    project = Project(Path(args.project))
    result = check_project(project)
    print(json.dumps(result, ensure_ascii=False, indent=1))


def cmd_render(args) -> None:
    project = Project(Path(args.project))
    result = check_project(project)
    if result["errors"] and not args.force:
        result["status"] = "needs_revision"
        result["message"] = "未导出：先修正 errors 再 render"
        print(json.dumps(result, ensure_ascii=False, indent=1))
        sys.exit(2)
    out = Path(args.output).resolve()
    if out.suffix.lower() != ".pptx":
        die("输出文件必须以 .pptx 结尾")
    if out.exists() and not args.overwrite:
        die(f"{out} 已存在；换一个文件名，或确认可覆盖后加 --overwrite")
    out.parent.mkdir(parents=True, exist_ok=True)
    platform = args.platform
    if platform == "auto":
        platform = "mac" if sys.platform == "darwin" else "win" if sys.platform.startswith("win") else "linux"
    try:
        count = Renderer(project, platform).render(out)
    except Exception as error:  # noqa: BLE001 - report the failing page/element instead of a bare trace
        die(f"渲染失败：{type(error).__name__}: {error}")
    from pptx import Presentation

    reopened = len(Presentation(str(out)).slides)
    print(json.dumps({
        "status": "exported" if reopened == count else "error",
        "output": str(out),
        "project": str(project.root),
        "pages": reopened,
        "fontPlatform": platform,
        "warnings": result["warnings"],
        "note": "已通过 python-pptx 重新打开校验；未在 PowerPoint/WPS 中做视觉验证。",
    }, ensure_ascii=False, indent=1))


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(prog="pptd.py", description="DSH PPT: PPTD 工程 → 可编辑 PPTX")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("templates", help="列出内置模板")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_templates)
    p = sub.add_parser("show", help="查看模板版式；--page 输出某页完整 PPTD")
    p.add_argument("template")
    p.add_argument("--lang", choices=["zh", "en"], default="zh")
    p.add_argument("--page", type=int)
    p.set_defaults(func=cmd_show)
    p = sub.add_parser("new", help="从模板复制一个新工程")
    p.add_argument("template")
    p.add_argument("project")
    p.add_argument("--lang", choices=["zh", "en"], default="zh")
    p.add_argument("--pages", help="选取并排序模板页，如 1,3,4,12")
    p.add_argument("--title")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_new)
    p = sub.add_parser("inspect", help="列出工程每页元素、位置、字号和文本容量")
    p.add_argument("project")
    p.add_argument("--page", type=int)
    p.add_argument("--all", action="store_true", help="也列出形状和线条")
    p.set_defaults(func=cmd_inspect)
    p = sub.add_parser("check", help="只读检查，返回 JSON")
    p.add_argument("project")
    p.set_defaults(func=cmd_check)
    p = sub.add_parser("render", help="检查并导出 PPTX")
    p.add_argument("project")
    p.add_argument("-o", "--output", required=True)
    p.add_argument("--platform", choices=["auto", "mac", "win", "linux"], default="auto", help="中文字体按哪个平台选择")
    p.add_argument("--overwrite", action="store_true")
    p.add_argument("--force", action="store_true", help="忽略 errors 强制导出（仅用于排查）")
    p.set_defaults(func=cmd_render)
    args = parser.parse_args(argv)
    try:
        import pptx  # noqa: F401
    except ImportError:
        if args.command in ("render",):
            die("缺少 python-pptx：用 load_workspace_dependencies 返回的 Python 运行本脚本")
    args.func(args)


if __name__ == "__main__":
    main()
